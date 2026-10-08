"""ATS Resume Checker - Streamlit app powered by Google Gemini Flash."""

import io
import json
import os

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = "gemini-2.5-flash"
MAX_RESUME_CHARS = 20000  # keeps prompts small and fast
MIN_RESUME_CHARS = 150  # below this, the file is probably scanned/empty

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume text provided and respond ONLY \
with JSON matching the requested schema.

Scoring guide (each section is 0-100):
- formatting: clean structure, standard section headings, consistent dates, no \
tables/columns/graphics that break ATS parsing, sensible length.
- keywords: relevant industry/role keywords and skills; if a job description is \
given, match against it.
- content: quantified achievements, strong action verbs, clarity, impact.
- completeness: contact info, summary, experience, education, skills present.

overall_score is 0-100 and should reflect the section scores realistically. Be \
honest and strict - do not inflate scores. Improvements must be specific and \
actionable, referencing the resume's actual content where possible."""

RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "overall_score": {"type": "integer"},
        "section_scores": {
            "type": "object",
            "properties": {
                "formatting": {"type": "integer"},
                "keywords": {"type": "integer"},
                "content": {"type": "integer"},
                "completeness": {"type": "integer"},
            },
            "required": ["formatting", "keywords", "content", "completeness"],
        },
        "summary": {"type": "string"},
        "strengths": {"type": "array", "items": {"type": "string"}},
        "improvements": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "priority": {"type": "string", "enum": ["High", "Medium", "Low"]},
                    "issue": {"type": "string"},
                    "suggestion": {"type": "string"},
                },
                "required": ["priority", "issue", "suggestion"],
            },
        },
        "missing_keywords": {"type": "array", "items": {"type": "string"}},
    },
    "required": [
        "overall_score",
        "section_scores",
        "summary",
        "strengths",
        "improvements",
        "missing_keywords",
    ],
}


# ----------------------------- Text extraction -----------------------------
def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
        except Exception:
            raise ValueError("Could not open this PDF. The file may be corrupted.")
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ------------------------------ Gemini analysis ----------------------------
def clamp(value, low=0, high=100) -> int:
    try:
        return max(low, min(high, int(round(float(value)))))
    except (TypeError, ValueError):
        return low


def parse_result(raw_text: str) -> dict:
    """Parse and sanitise the model's JSON so the UI never crashes."""
    cleaned = raw_text.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        if cleaned.lower().startswith("json"):
            cleaned = cleaned[4:]
    data = json.loads(cleaned.strip())

    sections = data.get("section_scores") or {}
    result = {
        "overall_score": clamp(data.get("overall_score")),
        "section_scores": {
            key: clamp(sections.get(key))
            for key in ("formatting", "keywords", "content", "completeness")
        },
        "summary": str(data.get("summary", "")),
        "strengths": [str(s) for s in (data.get("strengths") or [])],
        "improvements": [],
        "missing_keywords": [str(k) for k in (data.get("missing_keywords") or [])],
    }
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        result["improvements"].append(
            {
                "priority": priority,
                "issue": str(item.get("issue", "")),
                "suggestion": str(item.get("suggestion", "")),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    result["improvements"].sort(key=lambda i: order[i["priority"]])
    return result


def analyze_resume(api_key: str, resume_text: str, job_description: str = "") -> dict:
    client = genai.Client(api_key=api_key)
    prompt = f"RESUME TEXT:\n\"\"\"\n{resume_text[:MAX_RESUME_CHARS]}\n\"\"\"\n"
    if job_description.strip():
        prompt += (
            f"\nTARGET JOB DESCRIPTION:\n\"\"\"\n{job_description.strip()[:8000]}\n\"\"\"\n"
            "Score keyword match against this job description."
        )
    else:
        prompt += "\nNo job description was given; evaluate for general ATS-friendliness."

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            response_mime_type="application/json",
            response_schema=RESPONSE_SCHEMA,
            temperature=0.2,
        ),
    )
    if not response.text:
        raise RuntimeError("The model returned an empty response. Please try again.")
    return parse_result(response.text)


# ---------------------------------- UI -------------------------------------
def get_api_key() -> str:
    """Look for the key in Streamlit secrets, then env vars, then the sidebar."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets file locally
    env_key = os.environ.get("GEMINI_API_KEY", "")
    if env_key:
        return env_key
    return st.sidebar.text_input(
        "Gemini API key",
        type="password",
        help="Get a free key at https://aistudio.google.com/apikey",
    )


def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Excellent"
    if score >= 60:
        return "🟡 Good, needs work"
    return "🔴 Needs major improvement"


def render_result(result: dict) -> None:
    score = result["overall_score"]
    st.subheader("Your ATS Score")
    col1, col2 = st.columns([1, 2])
    col1.metric("Overall", f"{score} / 100")
    col2.markdown(f"### {score_label(score)}")
    st.progress(score / 100)
    if result["summary"]:
        st.write(result["summary"])

    st.subheader("Section breakdown")
    cols = st.columns(4)
    labels = {
        "formatting": "Formatting",
        "keywords": "Keywords",
        "content": "Content",
        "completeness": "Completeness",
    }
    for col, (key, label) in zip(cols, labels.items()):
        col.metric(label, result["section_scores"][key])

    if result["strengths"]:
        st.subheader("✅ Strengths")
        for s in result["strengths"]:
            st.markdown(f"- {s}")

    if result["improvements"]:
        st.subheader("🛠️ Suggested improvements")
        icons = {"High": "🔴", "Medium": "🟡", "Low": "🟢"}
        for item in result["improvements"]:
            with st.expander(f"{icons[item['priority']]} {item['priority']}: {item['issue']}"):
                st.write(item["suggestion"])

    if result["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_api_key()

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste the job description to check how well your resume matches it...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.getvalue(), uploaded.name)
            if len(text) < MIN_RESUME_CHARS:
                st.error(
                    "Could not read enough text from this file. If it is a scanned "
                    "image PDF, please upload a text-based PDF or a DOCX instead."
                )
                return
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, text, job_description)
        except json.JSONDecodeError:
            st.error("The AI returned an unreadable response. Please try again.")
            return
        except ValueError as exc:
            st.error(str(exc))
            return
        except Exception as exc:  # network, quota, invalid key, etc.
            st.error(f"Something went wrong: {exc}")
            return
        render_result(result)

    st.divider()
    st.caption("Scores are AI-generated estimates, not guarantees of how a real ATS will rank you.")


if __name__ == "__main__":
    main()
